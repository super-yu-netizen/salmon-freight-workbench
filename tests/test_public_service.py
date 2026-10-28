from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import unittest
from unittest.mock import patch

from openpyxl import load_workbook
from calculator import InputError, calculate
from public_service import (calculate_public, checked_snapshot, digest, parse_body,
                            parse_inquiry, validate_public)
from portable_documents import ExportError, export_json, export_inquiry, export_xlsx, export_pdf

ROOT = Path(__file__).resolve().parents[1]


class PublicServiceTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.catalog = json.loads((ROOT / 'data/catalog.json').read_text())
        cls.request = deepcopy(cls.catalog['default_request'])
        cls.request.update(supplier_ids=['DEMO-A', 'DEMO-B'], departure_date='2029-01-12')
        cls.result = calculate_public({'request': cls.request, 'catalog': cls.catalog}, cls.catalog)
        cls.snapshot = {'id': 'public-test', 'number': 'FE-20261028-PUBLIC-V1',
                        'group_id': 'FE-20261028-PUBLIC', 'revision': 1,
                        'created_at': '2026-10-28T03:00:00+00:00',
                        'request': cls.request, 'result': cls.result,
                        'catalog_snapshot': cls.catalog, 'catalog_hash': digest(cls.catalog),
                        'rule_version': cls.result['engine_version']}

    def test_same_core_amounts_and_no_catalog_mutation(self):
        before = deepcopy(self.catalog)
        expected = calculate(self.request, self.catalog)
        self.assertEqual(self.result['results'], expected['results'])
        self.assertEqual(self.result['catalog_hash'], digest(self.catalog))
        self.assertEqual(self.catalog, before)
        self.assertTrue(all(r['status'] == 'simulated' for r in self.result['results']))

    def test_per_visitor_catalog_is_used(self):
        changed = deepcopy(self.catalog)
        changed['version'] = 'CAT-VISITOR-TEST'
        out = calculate_public({'request': self.request, 'catalog': changed}, self.catalog)
        self.assertEqual(out['catalog_version'], 'CAT-VISITOR-TEST')
        self.assertNotEqual(self.catalog['version'], changed['version'])

    def test_validate_checks_contract_without_pricing_default_request(self):
        changed = deepcopy(self.catalog)
        changed['default_request']['departure_date'] = '2029-01-12'
        with patch('public_service.calculate', side_effect=AssertionError('must not price during validation')):
            self.assertEqual(validate_public({'catalog': changed})['catalog'], changed)

    def test_body_rejects_non_object_and_nonfinite(self):
        for raw in ('[]', 'null', '{"x": NaN}', '{"x": Infinity}', ''):
            with self.assertRaises(InputError):
                parse_body(raw)
        self.assertEqual(parse_body('{"text":"中文"}'), {'text': '中文'})

    def test_unknown_catalog_fails_before_calculating(self):
        with self.assertRaises(InputError):
            calculate_public({'request': self.request, 'catalog': {}}, self.catalog)

    def test_custom_catalog_text_match(self):
        changed = deepcopy(self.catalog)
        changed['products'][0]['name'] = '访客专用模拟产品'
        out = parse_inquiry('访客专用模拟产品；数量：1000；目的地新加坡', changed)
        self.assertEqual(out['draft']['product_id'], changed['products'][0]['id'])
        self.assertEqual(out['draft']['quantities'], [1000])

    def test_export_bounds_and_saved_snapshot(self):
        self.assertIs(checked_snapshot(self.snapshot), self.snapshot)
        malformed = deepcopy(self.snapshot)
        malformed['result']['results'] *= 5
        with self.assertRaises(InputError):
            checked_snapshot(malformed)
        with self.assertRaises(InputError):
            checked_snapshot({})

    def test_portable_xlsx_same_six_sheets_and_numeric_amounts(self):
        content = export_xlsx(self.snapshot)
        wb = load_workbook(BytesIO(content), data_only=True)
        self.assertEqual(wb.sheetnames, ['询价摘要', '包装明细', '计费依据', '费用明细', '情景对比', '来源与缺项'])
        exported = list(wb['情景对比'].iter_rows(min_row=5, values_only=True))
        self.assertEqual(len(exported), len(self.result['results']))
        for row, original in zip(exported, self.result['results']):
            self.assertEqual(row[7], float(original['budget_total']))
            self.assertEqual(row[8], float(original['per_set']))
            self.assertIn('模拟预估', row[10])
        all_text = '\n'.join(str(v) for row in wb['来源与缺项'].iter_rows(values_only=True) for v in row if v is not None)
        self.assertIn('不代表该日期市场报价', all_text)
        self.assertIn('SIM-DEMO-POLICY', all_text)

    def test_portable_formula_and_url_strings_stay_literal(self):
        changed = deepcopy(self.snapshot)
        changed['request']['metadata']['note'] = '=HYPERLINK("https://example.test","x")'
        changed['request']['metadata']['reference'] = '=1+1'
        wb = load_workbook(BytesIO(export_xlsx(changed)), data_only=False)
        for row in wb['询价摘要'].iter_rows():
            for cell in row:
                self.assertNotEqual(cell.data_type, 'f')

    def test_json_and_inquiry_keep_frozen_simulation(self):
        self.assertEqual(json.loads(export_json(self.snapshot)), self.snapshot)
        text = export_inquiry(self.snapshot).decode('utf-8-sig')
        self.assertIn('未发送', text)
        self.assertIn('不代表该日期市场报价', text)
        self.assertIn('2029-01-12', text)

    def test_pdf_requires_explicit_embeddable_font(self):
        with self.assertRaises(ExportError):
            export_pdf(self.snapshot)


if __name__ == '__main__':
    unittest.main()
