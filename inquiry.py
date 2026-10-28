"""Read the deliberately simple inquiry worksheet; never execute workbook code."""
import base64
from decimal import Decimal, InvalidOperation
from io import BytesIO
import json
import re
import zipfile


def read_inquiry_xlsx(encoded):
    if not isinstance(encoded, str):
        raise ValueError('请选择约定格式的询价Excel。')
    try:
        raw = base64.b64decode(encoded, validate=True)
    except (ValueError, TypeError):
        raise ValueError('Excel文件编码不正确。')
    if not 0 < len(raw) <= 3_000_000:
        raise ValueError('询价Excel应小于3MB。')
    try:
        with zipfile.ZipFile(BytesIO(raw)) as z:
            if len(z.infolist()) > 3000 or sum(x.file_size for x in z.infolist()) > 20_000_000:
                raise ValueError('文件过大，请使用简洁的询价模板。')
        import openpyxl
        book = openpyxl.load_workbook(BytesIO(raw), read_only=True, data_only=True, keep_links=False)
    except ValueError:
        raise
    except Exception:
        raise ValueError('无法读取，请使用下载的.xlsx询价模板。')
    try:
        if '询价输入' not in book.sheetnames:
            raise ValueError('请使用含“询价输入”工作表的约定模板。')
        sheet = book['询价输入']
        if (sheet.max_row is not None and sheet.max_row > 300) or (sheet.max_column is not None and sheet.max_column > 20):
            raise ValueError('询价模板内容超出约定范围。')
        aliases = {'产品ID': 'product_id', '产品/SKU': 'product_id', '产品SKU': 'product_id',
                   '产品': 'product_id', '数量': 'quantities', '数量档位': 'quantities',
                   '包装ID': 'packaging_id', '包装': 'packaging_id', '起运地': 'origin_id',
                   '目的地': 'destination_id', '服务': 'service_id', '运输方式': 'modes',
                   '货代': 'supplier_ids', '供应商': 'supplier_ids', '计划出运日期': 'departure_date',
                   '出运日期': 'departure_date', '预算汇率': 'fx_id', 'BU': 'metadata.bu',
                   '询价编号': 'metadata.reference', '备注': 'metadata.note', '演示估价': 'demo_estimation'}
        fields = {'product_id', 'quantities', 'packaging_id', 'origin_id', 'destination_id',
                  'service_id', 'modes', 'supplier_ids', 'departure_date', 'fx_id',
                  'metadata.bu', 'metadata.reference', 'metadata.note', 'demo_estimation'}
        draft = {}
        # Some valid writers omit the optional worksheet dimension metadata.
        # Read a bounded region instead of assuming max_row/max_column are set.
        for index, row in enumerate(sheet.iter_rows(max_row=301, max_col=21, values_only=True), start=1):
            if (index > 300 and any(x is not None for x in row)) or row[20] is not None:
                raise ValueError('询价模板内容超出约定范围。')
            if len(row) < 2 or row[0] is None or row[1] is None:
                continue
            field = str(row[0]).strip()
            field = aliases.get(field, field)
            if field not in fields:
                continue
            value = row[1]
            if hasattr(value, 'strftime'):
                value = value.strftime('%Y-%m-%d')
            if field in ('quantities', 'modes', 'supplier_ids'):
                try:
                    parsed = json.loads(str(value))
                except (ValueError, TypeError):
                    parsed = re.split(r'[,，、;；/\s]+', str(value).strip())
                if not isinstance(parsed, list):
                    parsed = [parsed]
                if field == 'quantities':
                    try:
                        checked = []
                        for x in parsed:
                            if not str(x).strip():
                                continue
                            number = Decimal(str(x))
                            if isinstance(x, bool) or not number.is_finite() or number != number.to_integral_value():
                                raise ValueError('not an integer')
                            checked.append(int(number))
                        parsed = checked
                    except (InvalidOperation, ValueError, TypeError):
                        raise ValueError('数量请填写整数数组，例如[1000,3000]。')
                    if not parsed or any(isinstance(x, bool) or x <= 0 or x > 1000000 for x in parsed):
                        raise ValueError('数量需要1至1000000的正整数。')
                if field == 'modes':
                    parsed = [{'空运': 'air', '海运': 'fcl', '海运整柜': 'fcl'}.get(x, x) for x in parsed]
                draft[field] = parsed
            elif field == 'demo_estimation':
                token = str(value).strip().lower()
                if token not in ('true', 'false', '开启', '关闭', '是', '否', '1', '0'):
                    raise ValueError('演示估价请填写true/false或开启/关闭。')
                draft[field] = token in ('true', '开启', '是', '1')
            elif field.startswith('metadata.'):
                draft.setdefault('metadata', {})[field.split('.')[1]] = str(value)
            else:
                draft[field] = str(value).strip()
        if not draft:
            raise ValueError('模板中没有可识别的询价字段。')
        return {'draft': draft, 'method': 'xlsx_template',
                'notes': ['已读取约定的“询价输入”表，请确认产品、数量和地点后测算。',
                          '只读取单元格值；不执行宏、外链或公式重算。']}
    finally:
        book.close()
